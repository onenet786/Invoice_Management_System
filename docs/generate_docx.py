import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_styled_document():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles Setup
    styles = doc.styles

    # Normal Style
    normal_style = styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Calibri'
    normal_font.size = Pt(11)
    normal_font.color.rgb = RGBColor(0x33, 0x41, 0x55) # Slate 700

    def set_cell_background(cell, hex_color):
        shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for margin, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
            node = OxmlElement(margin)
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def add_callout(text, title="NOTE", hex_border="4F46E5", hex_bg="EEF2FF"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, hex_bg)
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders_xml = (
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="none"/>'
            f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{hex_border}"/>'
            f'<w:bottom w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(parse_xml(borders_xml))

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r_title = p.add_run(f"[{title}] ")
        r_title.bold = True
        r_title.font.size = Pt(10)
        r_title.font.color.rgb = RGBColor(0x43, 0x38, 0xCA)

        r_text = p.add_run(text)
        r_text.font.size = Pt(10.5)
        r_text.font.color.rgb = RGBColor(0x1E, 0x1B, 0x4B)
        doc.add_paragraph()

    def add_screenshot_placeholder(caption):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=160, bottom=160, left=200, right=200)

        tcPr = cell._tc.get_or_add_tcPr()
        borders_xml = (
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
            f'<w:left w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
            f'<w:bottom w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
            f'<w:right w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(parse_xml(borders_xml))

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(8)
        
        r1 = p.add_run("[ SCREENSHOT PLACEHOLDER ]\n")
        r1.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        r2 = p.add_run(caption)
        r2.italic = True
        r2.font.size = Pt(10)
        r2.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        doc.add_paragraph()

    def add_header(title, level=1):
        h = doc.add_heading(level=level)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(title)
        if level == 1:
            h.paragraph_format.space_before = Pt(18)
            h.paragraph_format.space_after = Pt(6)
            run.font.size = Pt(18)
            run.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x1B, 0x4B) # Indigo 950
        elif level == 2:
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            run.font.size = Pt(14)
            run.bold = True
            run.font.color.rgb = RGBColor(0x31, 0x2E, 0x81) # Indigo 900
        elif level == 3:
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = RGBColor(0x43, 0x38, 0xCA) # Indigo 700

    # ------------------ COVER PAGE / TITLE ------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(40)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("INVOICEY")
    run_title.bold = True
    run_title.font.size = Pt(36)
    run_title.font.color.rgb = RGBColor(0x31, 0x2E, 0x81)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(20)
    run_sub = sub_p.add_run("Enterprise Invoice Management & Operations Manual")
    run_sub.font.size = Pt(16)
    run_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    meta_tbl = doc.add_table(rows=4, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Application:", "Invoicey Control Center v1.0.0"),
        ("Document Version:", "1.0 - Official Release"),
        ("Target Roles:", "Administrators, Managers, Viewers"),
        ("Platform Support:", "Windows Desktop, macOS, Linux, Android, iOS, Web")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_tbl.rows[idx]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(2.0)
        c1.width = Inches(4.5)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    doc.add_page_break()

    # ------------------ SECTION 1 ------------------
    add_header("1. Executive Summary & Overview", level=1)
    p = doc.add_paragraph(
        "Invoicey is an enterprise invoice management, quotation tracking, and billing operations system. "
        "Engineered for businesses managing equipment sales and specialized field services (including Solar PV installations "
        "and IT infrastructure deployment), the system combines offline-first local data resilience with automated cloud disaster "
        "recovery, native biometric access, and multi-channel client messaging."
    )
    p.paragraph_format.space_after = Pt(8)

    features = [
        ("Role-Based Security: ", "Full access segregation across Administrator, Manager, and Viewer accounts with real-time UI gating."),
        ("Biometric Authentication: ", "Integrated platform-level Fingerprint, Touch ID, and Face ID verification."),
        ("Offline-First Resilience: ", "Local storage persistence with auto-backup triggers to Google Drive upon any entity mutation."),
        ("Multi-Cloud Synchronization: ", "Cross-device REST API sync with Google authentication and optimistic revision concurrency."),
        ("Dynamic Invoicing Engine: ", "Tokenized sequential numbering, automated quotation-to-invoice transformation, OCR quote scanning, and 5 designer PDF templates.")
    ]
    for f_title, f_desc in features:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(4)
        r_bt = bp.add_run(f_title)
        r_bt.bold = True
        r_bt.font.color.rgb = RGBColor(0x1E, 0x1B, 0x4B)
        bp.add_run(f_desc)

    # ------------------ SECTION 2 ------------------
    add_header("2. First-Run Onboarding Gateway", level=1)
    doc.add_paragraph(
        "Upon launching Invoicey on a clean workstation, the user is presented with the Onboarding Gateway. "
        "All data remains localized on the host machine until remote synchronization is explicitly configured."
    )
    add_screenshot_placeholder("Onboarding Gateway screen displaying 'Set up my company' and 'Explore sample company' options")

    add_header("Option A: Set Up My Company (Production Deployment)", level=2)
    doc.add_paragraph("Follow these steps to initialize a fresh, production workspace:")
    steps_a = [
        "Select 'Set up my company'.",
        "Step 1 (Business Profile): Enter Legal Company Name, optional Tax ID / VAT / GST, select operating currency ($ / EUR / GBP / PKR / INR / AED), and provide official Office Address.",
        "Step 2 (Administrator Account): Provide primary Administrator Name, administrative Email Address, and Master Password.",
        "Click 'Complete Setup'. The company workspace is initialized, and you are redirected to the Login screen."
    ]
    for s in steps_a:
        doc.add_paragraph(s, style='List Number')

    add_header("Option B: Explore Sample Company (Demonstration Mode)", level=2)
    doc.add_paragraph(
        "Clicking 'Explore sample company' seeds the workspace with 'My Solar & IT Corp', 3 pre-configured user credentials "
        "(Admin: admin@invoice.com, Manager: manager@invoice.com, Viewer: viewer@invoice.com), 3 commercial clients, "
        "14 technical products across Solar and IT categories, and sample paid/overdue invoices."
    )

    # ------------------ SECTION 3 ------------------
    add_header("3. Authentication & Biometric Access", level=1)
    doc.add_paragraph(
        "Authentication validates user credentials against localized persistent storage and grants the corresponding role permissions."
    )
    add_screenshot_placeholder("Login screen featuring Email, Password fields, Sign In button, and Biometric fingerprint button")

    add_header("Step-by-Step Login Instructions:", level=2)
    steps_l = [
        "Launch Invoicey to open the Login screen.",
        "Enter your registered Email Address and Password.",
        "Click 'Sign In'. Upon verification, you are navigated to the main shell, displaying your assigned role badge in the top navigation bar.",
        "Biometric Sign-In: If biometric access is toggled ON in Settings, click 'Sign in with Biometrics' to trigger your device's fingerprint or facial recognition scanner for immediate password-free authentication."
    ]
    for s in steps_l:
        doc.add_paragraph(s, style='List Number')

    # ------------------ SECTION 4 ------------------
    add_header("4. Detailed Module Usage Guide", level=1)

    # Module 1
    add_header("Module 1: Dashboard & Cash Flow Analytics", level=2)
    doc.add_paragraph(
        "The Executive Dashboard computes real-time financial aggregates and displays cash flow trends across all invoices."
    )
    add_screenshot_placeholder("Dashboard view showing Total Revenue, Pending Payments, Overdue Amount, and Monthly Sales Bar Chart")
    doc.add_paragraph("Key Operational Steps:")
    d_steps = [
        "Financial KPI Review: Inspect Total Revenue (sum of all Paid invoices), Pending Payments (sum of Sent and Partially Paid invoices), and Overdue Amount (unsettled invoices past due date).",
        "Applying Filters: Use the Filter Controls Card to filter metrics by a single Client, specific Invoice Status, or Fiscal Year.",
        "Monthly Sales Trends: Hover over the monthly bar chart to analyze seasonal revenue distribution (Draft invoices are excluded from chart figures).",
        "Recent Invoices: Scroll to the bottom table to inspect recent transactions or click any row to navigate directly to that invoice's detail view."
    ]
    for s in d_steps:
        doc.add_paragraph(s, style='List Number')

    # Module 2
    add_header("Module 2: Client Directory Management", level=2)
    doc.add_paragraph("Maintains customer profiles, billing data, delivery logistics, and communication channels.")
    add_screenshot_placeholder("Client database screen displaying customer cards, search bar, and Add Client button")
    doc.add_paragraph("Adding a Client:")
    c_steps = [
        "Navigate to 'Clients' from the sidebar.",
        "Click '+ Add Client' in the upper right action bar.",
        "Populate Client Name, Accounting Email Address, International Phone Number, Billing Address, and optional Shipping Address.",
        "Click 'Save Client'. The record is instantly accessible in the Invoice Wizard."
    ]
    for s in c_steps:
        doc.add_paragraph(s, style='List Number')
    add_callout("Only Administrators and Managers can create, modify, or delete clients. Viewers have read-only inspection access.", "SECURITY RULE")

    # Module 3
    add_header("Module 3: Product Catalog & Automated SKU Engine", level=2)
    doc.add_paragraph("Central catalog for commercial stock items, hardware equipment, and professional services.")
    add_screenshot_placeholder("Product catalog view with Category tabs, SKU badges, unit prices, and Add Product dialog")
    p_steps = [
        "Navigate to 'Inventory' in the sidebar.",
        "Browse items by Category Tabs ('All', 'Solar', 'IT', 'Hardware', 'Services').",
        "Click '+ Add Product'.",
        "Enter Item Name (e.g., 'Commercial Solar Inverter 25kW') and Category.",
        "Click 'Auto Generate' next to SKU. Invoicey analyzes the title and category to generate an industry-standard sequential code (e.g., 'SOL-COMSOLIN-001').",
        "Input Unit Price and Tax Rate (%), then click 'Save Product'."
    ]
    for s in p_steps:
        doc.add_paragraph(s, style='List Number')

    # Module 4
    add_header("Module 4: 3-Step Invoice Creation Wizard", level=2)
    doc.add_paragraph("The primary commercial engine for building validated invoices and formal quotations.")
    add_screenshot_placeholder("3-Step Invoice Wizard displaying Step 1 Client Selection, Step 2 Line Items, and Step 3 Review")
    w_steps = [
        "Click 'Create Invoice' (or 'Create Quote') on the Invoices screen.",
        "Step 1 (Client Details): Verify the auto-generated sequential number, select the Target Client from the dropdown, and establish Issue and Due dates. Click 'Continue to Line Items'.",
        "Step 2 (Line Items): Select a catalog item from the dropdown. Unit price and tax rate populate automatically. Enter quantity, adjust line pricing if negotiated, and click 'Add Item to Invoice'. Repeat for additional products. Click 'Continue to Review'.",
        "Step 3 (Review & Issue): Verify Subtotal, Tax Total, and Grand Total. Select the initial status ('Draft', 'Sent', or 'Paid'), append customer terms or banking notes, and click 'Issue Invoice'."
    ]
    for s in w_steps:
        doc.add_paragraph(s, style='List Number')

    # Module 5
    add_header("Module 5: OCR Quotation Scanning", level=2)
    doc.add_paragraph(
        "Imports paper site surveys or handwritten contractor estimates directly into structured digital invoice items with zero manual typing."
    )
    add_screenshot_placeholder("Scan Quotation Dialog showing OCR text extraction and detected line items preview")
    ocr_steps = [
        "On the Invoices screen, click 'Scan Quote' in the header action bar.",
        "Select a scanned quote template (e.g., Solar System Layout Quote or IT Hardware Upgrade Quote).",
        "Inspect the detected handwritten notes and extracted line item table.",
        "Click 'Import Scanned Items'. The system opens the Invoice Wizard with client, items, quantities, and pricing fully pre-loaded."
    ]
    for s in ocr_steps:
        doc.add_paragraph(s, style='List Number')

    # Module 6
    add_header("Module 6: Quote-to-Invoice 1-Click Conversion", level=2)
    doc.add_paragraph(
        "Converts approved quotations into active tax invoices without re-entering data, maintaining full audit trail linkages."
    )
    add_screenshot_placeholder("Invoice Detail screen showing Quote document and 'Convert to Invoice' action button")
    q_steps = [
        "Open an existing Quote from the Invoices directory.",
        "Click 'Convert to Invoice' in the top action bar.",
        "The system generates a new sequential invoice number, sets status to Draft, copies all items, and stamps the original quote with the converted invoice ID to prevent duplicate conversions."
    ]
    for s in q_steps:
        doc.add_paragraph(s, style='List Number')

    # Module 7
    add_header("Module 7: Document Delivery & Sharing (PDF, WhatsApp, Email)", level=2)
    doc.add_paragraph("Delivers invoices through standard printing and modern instant messaging channels.")
    add_screenshot_placeholder("Document sharing modal showing WhatsApp Text, WhatsApp PDF, and Email buttons")
    m_steps = [
        "PDF Generation: On any invoice screen, click 'View PDF' to render a vector-crisp document using the active company template. Print or save directly from the viewer.",
        "WhatsApp Text Summary: Click 'WhatsApp' -> 'Send Text Summary' to launch WhatsApp with a formatted invoice breakdown pre-addressed to the client's phone number.",
        "WhatsApp PDF Dispatch: Click 'WhatsApp' -> 'Share PDF Document' to render the PDF and open the native system share sheet.",
        "Native Email Dispatch: Click 'Email' to trigger the system email client pre-populated with recipient address, subject line, and invoice breakdown. (Draft invoices automatically advance to Sent status)."
    ]
    for s in m_steps:
        doc.add_paragraph(s, style='List Number')

    # Module 8
    add_header("Module 8: Settings, Role Sandbox & Cloud Disaster Recovery", level=2)
    doc.add_paragraph("Centralized administrative control center for legal configuration, templates, and backups.")
    add_screenshot_placeholder("Settings screen showing Company Profile, Number Format, PDF Templates, and Cloud Sync")
    cfg_items = [
        ("Company Profile: ", "Configure legal name, tax identification, operating currency, address, and upload company logo (PNG/JPG up to 5 MB)."),
        ("Invoice Number Format: ", "Define sequential format tokens using {YYYY}, {YY}, {MM}, and exactly one sequence placeholder like {NNNN}."),
        ("Designer PDF Templates: ", "Switch between Classic, Modern, Minimal, Corporate, and Elegant templates with live modal preview."),
        ("Testing Sandbox: ", "Switch active session role between Admin, Manager, and Viewer to audit permission restrictions in real-time."),
        ("Google Drive Backup: ", "Link Google account, enable Auto-Backup upon every mutation, and restore from a rolling ledger of 10 historical snapshots."),
        ("Remote Cloud Sync: ", "Connect to HTTPS REST backend with Google ID Token authentication and optimistic concurrency conflict (HTTP 409) protection.")
    ]
    for c_lbl, c_val in cfg_items:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(4)
        r = bp.add_run(c_lbl)
        r.bold = True
        r.font.color.rgb = RGBColor(0x1E, 0x1B, 0x4B)
        bp.add_run(c_val)

    # ------------------ SECTION 5 ------------------
    add_header("5. Role-Based Operations Guide", level=1)
    doc.add_paragraph("Invoicey implements strict Role-Based Access Control (RBAC). The operational matrix is detailed below:")

    role_tbl = doc.add_table(rows=12, cols=4)
    role_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Feature / Operation", "Admin", "Manager", "Viewer"]
    for idx, h in enumerate(headers):
        c = role_tbl.rows[0].cells[idx]
        set_cell_background(c, "1E1B4B")
        set_cell_margins(c, 100, 100, 100, 100)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    matrix_data = [
        ("View Dashboard & Analytics", "YES", "YES", "YES"),
        ("Browse Clients, Products, Invoices", "YES", "YES", "YES"),
        ("View & Download PDF Documents", "YES", "YES", "YES"),
        ("Create & Edit Invoices / Quotes", "YES", "YES", "NO"),
        ("Delete Invoices", "YES", "YES", "NO"),
        ("Add, Edit, Delete Clients", "YES", "YES", "NO"),
        ("Add, Edit, Delete Products", "YES", "YES", "NO"),
        ("Trigger WhatsApp & Email Reminders", "YES", "YES", "NO"),
        ("Modify Company Profile & Logo", "YES", "NO", "NO"),
        ("Edit Invoice Numbering Format", "YES", "NO", "NO"),
        ("Configure Cloud Sync & Drive Snapshots", "YES", "NO", "NO")
    ]
    for r_idx, row_data in enumerate(matrix_data, start=1):
        row = role_tbl.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if val == "YES":
                r.bold = True
                r.font.color.rgb = RGBColor(0x16, 0x65, 0x34) # Green 800
            elif val == "NO":
                r.bold = True
                r.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B) # Red 800
            else:
                r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # ------------------ SECTION 6 ------------------
    add_header("6. Troubleshooting & Problem Resolution", level=1)
    issues = [
        ("Access Denied: Viewers cannot create or edit invoices",
         "The current logged-in account has the Viewer role. Log out and sign in with a Manager or Admin account. If evaluating in demonstration mode, open Settings -> Testing Sandbox and toggle to Admin."),
        ("Cloud data changed on another device. Download it before uploading again.",
         "Optimistic concurrency locking conflict (HTTP 409). Another user synchronized changes to the cloud server, incrementing the workspace revision counter. Open Settings -> Remote Cloud Sync, click 'Download Workspace' to merge remote data, then proceed with your upload."),
        ("Biometric verification failed / Not available",
         "The host workstation lacks biometric hardware or no fingerprints/facial profiles are enrolled in the operating system. Use email/password authentication instead."),
        ("Error: Selected client does not have a valid email address configured",
         "The client profile attached to the invoice is missing a valid email address. Open Clients, edit the client card, enter an RFC-compliant email, save, and re-launch email dispatch."),
        ("Logo must be smaller than 5 MB",
         "The imported image exceeds the 5 MB threshold. Compress the PNG or JPG using standard image utilities prior to upload.")
    ]
    for issue_title, issue_res in issues:
        add_header(issue_title, level=3)
        doc.add_paragraph(issue_res)

    # ------------------ SECTION 7 ------------------
    add_header("7. Frequently Asked Questions (FAQs)", level=1)
    faqs = [
        ("Does Invoicey require an internet connection to issue invoices?",
         "No. Invoicey operates on an offline-first architecture. All clients, inventory items, and invoices are written to local persistent storage. Internet connectivity is only required for remote cloud sync and email/WhatsApp dispatch."),
        ("How are overdue invoices detected?",
         "The system checks all invoices whenever the application loads. Any invoice in Sent or Partially Paid status whose dueDate is earlier than the current system timestamp is automatically reclassified as Overdue."),
        ("What happens to a Quote when converted to an Invoice?",
         "The original Quote is preserved for historical audit trails. Its convertedInvoiceId property is populated with the ID of the newly spawned invoice, preventing duplicate conversions while preserving original quotation terms."),
        ("Where are generated PDF invoices stored?",
         "The built-in PDF viewer interfaces with your operating system's native printing framework, allowing you to save files to any local directory, cloud drive, or physical printer.")
    ]
    for q, a in faqs:
        qp = doc.add_paragraph()
        qr = qp.add_run(f"Q: {q}")
        qr.bold = True
        qr.font.color.rgb = RGBColor(0x31, 0x2E, 0x81)
        ap = doc.add_paragraph()
        ar = ap.add_run(f"A: {a}")
        ar.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        ap.paragraph_format.space_after = Pt(8)

    output_path = os.path.join(os.path.dirname(__file__), "USER_GUIDE.docx")
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    create_styled_document()
